"""Extrae SOLO equipos (no cables, accesorios, tubería, válvulas, instrumentos ni materiales) de los tracking
ECU17049 (M3), ECU21212 (M4-E) y ECU24091 (M5) -> data_trk.json, cruzado con la ruta del taladro (RyC)."""
import sys, re, json, os, glob, datetime, collections, warnings
warnings.simplefilter("ignore")
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
D = r"C:\OD\OneDrive - ApplusGlobal\Documentos\Edgar_Applus\Tracking"
FILES = [("M3", glob.glob(D + r"\Exteriores Elkin\ECU17049_*.xlsx")), ("M4", glob.glob(D + r"\Estaciones Paola\ECU21212_*.xlsx")),
         ("M5", glob.glob(D + r"\Exteriores Elkin\ECU24091_*.xlsx"))]
TAL = r"C:\Users\edgar.suarez\Dashboard\Timeline_Taladro\Propuesta_v2\data_v2.json"
HOY = datetime.date.today()

def iso(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat() if v.year > 1990 else None
    if isinstance(v, datetime.date): return v.isoformat()
    if isinstance(v, str):
        m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", v)
        if m: return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return None
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0
def s(v): return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""

# --- qué es equipo y en qué categoría cae
EXCL_TIPO = re.compile(r"FOTOVOLT|ILUMIN|LUMINAR|PANEL SOLAR|DISPLAY|TARJETA|RELAY|REPUEST|BOQUILLA|FILTRO EN Y|FILTRO TEMPORAL|ESCOTILLA|TRAMPA|MANGUERA|ABRAZADERA|CONDUIT|CONDULETA|CONECTOR|PERNO|REDUCCI|RIEL|ACCESORIO|BANDEJA|LICENCIA|PATCHCORD|PIG TAIL|SFP|KIT DE AMPLIACION|SERVICIOS|CONFIGURACI|PRUEBAS|INTEGRACI|ACTUALIZACI"
                       r"|^DPS$|PROTECTOR DE AVES|CLAMP|^LAMP|TERMINAL|^CABLE|CIRCUIT|^MODULE|^PLATE|^SWITCH$|CONVERTIDOR|MASTIL|O-RING|PARTS KIT|SPARE PARTS"
                       r"|^CAJAS?\b|CAJA DE COMUNICACIONES|^FILTRO$|STRAINER|AMPLIACIÓN EN TABLEROS", re.I)
CAT = [
    ("Transformadores", lambda t, tm: tm == "TRANSFORMADORES"),
    ("Skids VFD/SUT y E-House", lambda t, tm: tm == "EQ. ELÉCTRICO SKID"),
    ("Tableros y celdas", lambda t, tm: tm in ("TABLEROS Y CEL. ELÉC", "TABLEROS DE CONTROL") or (tm == "MAT. INST. Y CONTROL" and re.search(r"GABINETE|CENTRO MANIOBRA|DATA CENTER|TABLERO", t))),
    ("Variadores", lambda t, tm: tm == "VARIADORES"),
    ("Otros eléctricos", lambda t, tm: tm == "OTROS EQUIPOS ELÉCTRICOS"),
    ("Bombas y rotativos", lambda t, tm: tm == "EQ. ROTATIVO"),
    ("Recipientes y estáticos", lambda t, tm: tm in ("RECIPIENTE A PRESIÓN", "EQ. ESTÁTICO")),
]
def categoria(tipo, tm):
    if EXCL_TIPO.search(tipo): return None
    for name, f in CAT:
        if f(tipo.upper(), tm): return name
    return None
def zona(mod, sub):
    u = sub.upper()
    if mod in ("M3", "M5"): return "Exteriores"
    if re.search(r"CLUSTER|CLÚSTER|RAMAL|APPRAISAL", u): return "Exteriores"
    return "Estaciones"
def loc(*vals):
    for v in vals:
        m = re.search(r"\b(CASE|CUMA|EMBR|WDCS|CMO)[\s\-_]*0*(\d{1,4})", str(v or ""), re.I)
        if m: return f"{m.group(1).upper()}-{int(m.group(2)):04d}"
    return None

# --- ruta del taladro: por clúster, llegada del taladro y 1er completamiento
tal = json.load(open(TAL, encoding="utf-8"))
TC = {}
for x in tal["items"]:
    c = loc(x["cluster"]) or x["cluster"]
    t = TC.setdefault(c, dict(taladro=set(), ini=None, comp1=None, fin=None))
    t["taladro"].add(x["taladro"].split(" - ")[0])
    if x["ini"]: t["ini"] = min(t["ini"] or x["ini"], x["ini"])
    if x["fin"]: t["fin"] = max(t["fin"] or x["fin"], x["fin"])
    if x["evento"] == "COMPLETAMIENTO" and x["fin"]: t["comp1"] = min(t["comp1"] or x["fin"], x["fin"])

items, fuentes = [], {}
for mod, fs in FILES:
    if not fs: continue
    f = max(fs, key=os.path.getmtime); fuentes[mod] = dict(archivo=os.path.basename(f), fecha=datetime.date.fromtimestamp(os.path.getmtime(f)).isoformat())
    ws = load_workbook(f, read_only=True, data_only=True)["GENERAL"]
    it = ws.iter_rows(values_only=True); H = [s(h) for h in next(it)]
    ix = {h: i for i, h in enumerate(H)}
    g = lambda r, k: r[ix[k]] if k in ix and ix[k] < len(r) else None
    for r in it:
        if not any(r): continue
        tm = s(g(r, "Tipología Material")).upper(); tipo = s(g(r, "Tipo de material"))
        cat = categoria(tipo, tm)
        if not cat: continue
        est = s(g(r, "Estado del Proceso de Aprovisionamiento"))
        if est.lower().startswith("cancel") or s(g(r, "ESTADO DE LA ODB")).upper() == "CANCELADA": continue
        cant = num(g(r, "Cantidad por ítem"))
        if cant <= 0: continue
        sub = s(g(r, "Sub Proyecto"))
        a1, a2, tag = s(g(r, "Área de instalación 1 (Sistema)")), s(g(r, "Área de instalación 2 (Sub-Sistema)")), s(g(r, "TAG"))
        rec, pend = num(g(r, "Cantidad recibida en Sitio")), num(g(r, "Cantidad Pendiente por Recibir"))
        eta = iso(g(r, "Fecha Proyectada de Recibo en sitio o Fecha Real Recibido en Sitio"))
        fsap = iso(g(r, "Fecha de Recibido en SAP"))
        recibido = est.lower().startswith("recib") or (rec >= cant and rec > 0) or (pend == 0 and fsap is not None)
        if recibido: pend = 0
        elif pend <= 0: pend = max(cant - rec, 0) or cant
        z = zona(mod, sub); l = loc(a1, tag, a2)
        t = TC.get(l, {}) if l else {}
        freq_proj = iso(g(r, "Fecha requerida por el proyecto"))
        # fecha requerida en sitio: para clústeres con taladro, la llegada del taladro (ini); si no, la fecha requerida del proyecto
        req, req_src = (t.get("ini"), "Llegada del taladro") if t.get("ini") else (freq_proj, "Fecha requerida proyecto" if freq_proj else None)
        items.append(dict(
            mod=mod, zona=z, sub=sub, cat=cat, tipologia=tm.title(), tipo=tipo, desc=s(g(r, "Descripción del ítem"))[:220], tag=tag,
            loc=l or (a1[:40] if a1 else "—"), a1=a1, a2=a2, cant=cant, und=s(g(r, "Unidad del ítem")), est=est or "—",
            prov=s(g(r, "Proveedor")) or "—", odb=s(g(r, "No. Orden de Compra")), item_oc=s(g(r, "Ítem de la OC")),
            f_oc=iso(g(r, "Fecha Entrega OC")), est_plan=s(g(r, "Estado Planeación ODB")), insp=s(g(r, "Estado de Inspección del Ítem")),
            f_lib=iso(g(r, "Fecha final de liberación del ítem")), eta=eta, f_sap=fsap, rec=rec, pend=pend, est_odb=s(g(r, "ESTADO DE LA ODB")),
            obs=s(g(r, "Observaciones Proceso de Compra"))[:300], obs_bod=s(g(r, "Observaciones Llegada Material a Bodegas"))[:300],
            f_req_proj=freq_proj, req=req, req_src=req_src, taladro=", ".join(sorted(t.get("taladro", []))), tal_ini=t.get("ini"), comp1=t.get("comp1"),
            recibido=recibido, mr=s(g(r, "Código de MR"))))

# --- estado de llegada y alertas
def dd(a, b): return (datetime.date.fromisoformat(a) - datetime.date.fromisoformat(b)).days
for x in items:
    al = []
    if x["recibido"]:
        x["lleg"] = "Recibido"
    else:
        if not x["odb"] and not x["eta"]: x["lleg"] = "Sin compra"; al.append("Sin ODB · " + x["est"].lower())
        elif not x["eta"]: x["lleg"] = "Sin fecha"; al.append("Con ODB pero sin fecha de llegada")
        elif x["eta"] < HOY.isoformat(): x["lleg"] = "ETA vencida"; al.append(f"ETA vencida hace {dd(HOY.isoformat(), x['eta'])} días")
        elif dd(x["eta"], HOY.isoformat()) <= 30: x["lleg"] = "Llega ≤30 días"
        elif dd(x["eta"], HOY.isoformat()) <= 90: x["lleg"] = "Llega 31–90 días"
        else: x["lleg"] = "Llega >90 días"
        if x["eta"] and x["req"] and x["eta"] > x["req"]: al.append(f"Llega {dd(x['eta'], x['req'])} días después de lo requerido ({x['req_src'].lower()})")
        if not x["eta"] and x["tal_ini"] and 0 <= dd(x["tal_ini"], HOY.isoformat()) <= 120:
            al.append(f"El taladro llega en {dd(x['tal_ini'], HOY.isoformat())} días y el equipo no tiene fecha")
        if x["eta"] and x["odb"] == "" : al.append("Sin ODB")
    x["alertas"] = al
    x["holgura"] = dd(x["req"], x["eta"]) if (x["eta"] and x["req"] and not x["recibido"]) else None
items.sort(key=lambda x: (x["recibido"], x["eta"] or "9999", x["zona"], x["loc"]))
tal_out = {c: dict(taladro=sorted(t["taladro"]), ini=t["ini"], comp1=t["comp1"], fin=t["fin"]) for c, t in TC.items()}
json.dump(dict(hoy=HOY.isoformat(), fuentes=fuentes, corte_ryc=tal.get("corte"), items=items, taladro=tal_out),
          open(os.path.join(HERE, "data_trk.json"), "w", encoding="utf-8"), ensure_ascii=False)
print("fuentes:", fuentes)
print("equipos (líneas):", len(items), "| unidades:", sum(x["cant"] for x in items))
for k in ("zona", "mod", "cat", "lleg", "est"):
    print(k, collections.Counter(x[k] for x in items).most_common())
pend = [x for x in items if not x["recibido"]]
print("pendientes:", len(pend), "| con alerta de llegar tarde:", sum(1 for x in pend if any("después" in a for a in x["alertas"])))
print("ubicaciones pendientes:", collections.Counter(x["loc"] for x in pend).most_common(25))
for x in pend[:25]:
    print(f"  {x['eta']} | {x['zona'][:4]} {x['mod']} | {x['loc'][:12]:12} | {x['cat'][:16]:16} | {x['tipo'][:30]:30} | {x['prov'][:22]:22} | ODB {x['odb']} | req {x['req']} | {x['alertas']}")

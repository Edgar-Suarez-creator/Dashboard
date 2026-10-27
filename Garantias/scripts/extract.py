import openpyxl, datetime as dt, re, json, sys
def clean(v):
    if v is None: return None
    if isinstance(v,str):
        v=re.sub(r'[ \t]+',' ',v.replace('\xa0',' ')).strip()
        return v or None
    return v
def d(v):
    v=clean(v)
    if isinstance(v,dt.datetime): return v.date().isoformat() if v.year>1900 else None
    return None
def ds(v):
    v=clean(v)
    if isinstance(v,dt.datetime): return d(v)
    m=re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})',str(v or ''))
    if m:
        try: return dt.date(int(m[3]),int(m[2]),int(m[1])).isoformat()
        except ValueError: pass
    return s(v)
def s(v):
    v=clean(v)
    if v is None: return None
    if isinstance(v,float) and v.is_integer(): v=int(v)
    if isinstance(v,dt.datetime): return v.date().isoformat()
    return str(v)
def disc(v):
    v=(clean(v) or '').lower()
    if v.startswith('el'): return 'Eléctrica'
    if 'i&c' in v: return 'I&C'
    if v.startswith('mec'): return 'Mecánica'
    if v.startswith('civ'): return 'Civil'
    return '(sin dato)' if not v else v
def edate(x,m):
    y,mo=x.year+(x.month-1+m)//12,(x.month-1+m)%12+1
    import calendar
    return dt.date(y,mo,min(x.day,calendar.monthrange(y,mo)[1]))
def extract(path,today):
    ws=openpyxl.load_workbook(path,data_only=True)['MATRIZ']
    T=dt.date.fromisoformat(today); out=[]
    for r in range(2,ws.max_row+1):
        c=lambda i: ws.cell(r,i).value
        if all(clean(c(i)) is None for i in range(1,33)): continue
        proy=s(c(2)); rep=s(c(23))
        fa,fc=d(c(24)),d(c(25))
        est=(s(c(29)) or 'SIN CLASIFICAR').upper()
        if est not in ('ABIERTA','CERRADA'): est='SIN CLASIFICAR'
        F,G,H=[clean(c(i)) for i in (6,7,8)]
        F=F.date() if isinstance(F,dt.datetime) and F.year>1900 else None
        G=G.date() if isinstance(G,dt.datetime) else None
        H=H.date() if isinstance(H,dt.datetime) else None
        lim=min([x for x in (edate(F,18) if F else None, edate(G,12) if G else None) if x] or [None]) if (F or G) else None
        vig=None
        if H and lim: vig='VIGENTE' if H<=lim else 'EXPIRADA'
        elif clean(c(9)) and not str(clean(c(9))).startswith('='): vig=s(c(9))
        atr=(lim-T).days if (est=='ABIERTA' and lim) else None
        dur=(dt.date.fromisoformat(fc)-dt.date.fromisoformat(fa)).days if fa and fc else None
        dias=dur if est=='CERRADA' else ((T-dt.date.fromisoformat(fa)).days if fa else None)
        ag=None
        if est=='ABIERTA' and dias is not None:
            ag='0–90 días' if dias<=90 else '91–180 días' if dias<=180 else '181–365 días' if dias<=365 else 'Más de 365 días'
        vip=re.sub(r'\s+','',(rep or '').split('\n')[0])
        bit=s(c(27))
        dates=re.findall(r'^\s*(\d{1,2})/(\d{1,2})/(\d{4})',bit or '')
        ub=None
        if dates:
            dl=[]
            for a,b,y in dates:
                try: dl.append(dt.date(int(y),int(b),int(a)))
                except: pass
            if dl: ub=max(dl).isoformat()
        nivel=s(c(18)); nivel='N/a' if nivel and nivel.upper()=='N/A' else nivel
        grupo=s(c(17)); grupo=grupo.capitalize() if grupo else None
        out.append(dict(fila=r,linea=c(1),mod=('M'+proy[-1]) if proy else None,proyecto=proy,proveedor=s(c(3)),
          odb=s(c(4)),fecha_odb=d(c(5)),fecha_entrega=ds(c(6)),aceptacion_hc1=ds(c(7)),fecha_notif=ds(c(8)),
          notif_txt=s(c(8)),vigencia=vig,atraso=atr,item_odb=s(c(11)),umb=s(c(12)),cant=clean(c(13)),
          cod_material=s(c(14)),tag=s(c(15)),tipo=s(c(16)) or '(sin dato)',grupo=grupo,nivel=nivel,disc=disc(c(19)),
          sistema=s(c(20)),subsistema=s(c(21)),ubic=s(c(22)),reporte=rep,fecha_apertura=fa,fecha_cierre=fc,
          descripcion=s(c(26)),bitacora=bit,seguimiento=d(c(28)),estado=est,cerrado_por=s(c(30)),contacto=s(c(31)),
          obs=s(c(32)),mes_apertura=fa[:7] if fa else None,mes_cierre=fc[:7] if fc else None,duracion=dur,dias=dias,
          aging=ag,vip=vip,key=f"{s(c(4))}|{vip}" if vip else f"L{c(1)}",ult_bitacora=ub))
    return out
if __name__=='__main__':
    json.dump(extract(sys.argv[1],sys.argv[2]),open(sys.argv[3],'w'),ensure_ascii=False,indent=1)

"""Dashboard Asistencias Técnicas · Vendors CSE: lee Control_Asistencias_Tecnicas_CSE.xlsx (Control, Calendario, Contratos) -> data_at.json."""
import os, re, sys, json, shutil, tempfile, datetime, collections, warnings
warnings.simplefilter("ignore")
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
MAIN = r"C:\OD\ApplusGlobal\Javier Alfonso Cruz Pinto - Planeación APPLUS\11. Plan de Vendor CoC1 M4\00. Dashboard\Control_Asistencias_Tecnicas_CSE.xlsx"
LOCAL = r"C:\OD\OneDrive - ApplusGlobal\Documentos\Edgar_Applus\Vendor\Control_Asistencias_Tecnicas_CSE.xlsx"
SRC = sys.argv[1] if len(sys.argv) > 1 else MAIN
t = os.path.join(tempfile.gettempdir(), "_ctrl_at_dash.xlsx")
try: shutil.copyfile(SRC, t)
except PermissionError: shutil.copyfile(LOCAL, t); SRC = LOCAL
wb = load_workbook(t, data_only=True)
MES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12}
def iso(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime("%Y-%m-%d")
    return None
def num(v):
    try: return float(v)
    except Exception: return None
def txt(v):
    if v is None: return ""
    s = str(v).strip()
    return "" if s in ("None", "—") else s
# ---------------- Control
c = wb["Control"]
H = {str(c.cell(4, j).value).strip(): j for j in range(1, 60) if c.cell(4, j).value}
def col(name):
    for k, j in H.items():
        if k.startswith(name): return j
    raise KeyError(name)
CHK = ["Solicitud de creación", "ODS creada", "ODS suscrita", "Pólizas aprobadas", "KOM", "RSPA", "Acta de inicio", "Equipo en sitio", "HSE / ingreso"]
CHK_S = ["Solicitud", "ODS creada", "Suscrita", "Pólizas", "KOM", "RSPA", "Acta inicio", "Equipo sitio", "HSE"]
ats = []; hitos = []
for r in range(5, 160):
    at = txt(c.cell(r, 1).value); eq = txt(c.cell(r, col("Equipo / alcance")).value)
    if not at or not eq: continue
    g = lambda n: c.cell(r, col(n)).value
    obs = txt(g("Observaciones"))
    x = dict(id=at, fr=txt(g("Frente")), mod=txt(g("Módulo")), ubic=txt(g("Ubicación")), eq=eq, ven=txt(g("Vendor")), ods=txt(g("ODS")), marco=txt(g("Contrato marco")),
             vig=iso(g("Vigencia ODS")), chk=[txt(g(k)) for k in CHK], dias=num(g("Días contratados")), freq=iso(g("Fecha requerida")), cons=num(g("Avance construcción")),
             odb=txt(g("ODB del equipo")), lleg=iso(g("Llegada del equipo")), esteq=txt(g("Estado del equipo")), dpend=num(g("Días programados pendientes")) or 0,
             dejec=num(g("Días ejecutados")) or 0, prox=iso(g("Próximo día programado")), ult=iso(g("Último día")), infp=num(g("Informes diarios pendientes")) or 0,
             pas=num(g("Días pasados sin registrar")) or 0, inff=txt(g("Informe final recibido")), acta=txt(g("Acta de cierre")), etapa=txt(g("Etapa")),
             falta=txt(g("¿Qué falta?")), sem=int(num(g("Semáforo")) or 0), mot=txt(g("Motivo de la alerta")), resp=txt(g("Responsable")), obs=obs[:3000])
    ats.append(x)
    # bitácora: notas con fecha al inicio ("24-oct: ...", "08-oct: ...") separadas por salto de línea o " | "
    for part in re.split(r"\n|\s\|\s", obs):
        m = re.match(r"\s*(\d{1,2})[-/ ](ene|feb|mar|abr|may|jun|jul|ago|sep|oct|nov|dic)[a-z]*\.?(?:[-/ ](\d{2,4}))?\s*[:·-]\s*(.+)", part, re.I)
        if not m: continue
        y = int(m.group(3)) if m.group(3) else 2026
        y = y + 2000 if y < 100 else y
        try: d = datetime.date(y, MES[m.group(2).lower()[:3]], int(m.group(1)))
        except ValueError: continue
        hitos.append(dict(at=at, ven=x["ven"], ubic=x["ubic"], f=d.isoformat(), t=m.group(4).strip()[:600]))
# ---------------- Calendario
k = wb["Calendario"]
days = [(j, iso(k.cell(4, j).value)) for j in range(14, 240) if iso(k.cell(4, j).value)]
cal = {}
for r in range(5, 160):
    at = txt(k.cell(r, 1).value)
    if not at: continue
    marks = [[d, str(k.cell(r, j).value).strip().upper()] for j, d in days if k.cell(r, j).value not in (None, "")]
    if marks: cal[at] = marks
# ---------------- Contratos
kc = wb["Contratos"]
HC = {str(kc.cell(4, j).value).strip(): j for j in range(1, 30) if kc.cell(4, j).value}
contr = []
for r in range(5, 120):
    ods = txt(kc.cell(r, HC["ODS"]).value)
    if not ods: continue
    gc = lambda n: kc.cell(r, HC[n]).value
    contr.append(dict(ven=txt(gc("Vendor")), marco=txt(gc("Contrato marco")), ods=ods, fr=txt(gc("Frente")), adm=txt(gc("Administrador ECP")), obj=txt(gc("Objeto / alcance")),
                      v0=num(gc("Valor inicial (COP)")), v1=num(gc("Valor actual (COP)")), ini=iso(gc("Vigencia inicio")), fin=iso(gc("Vigencia fin")),
                      dc=num(gc("Días contratados")), dant=num(gc("Días ejecutados antes del control")), dej=num(gc("Días ejecutados en el control")),
                      dpro=num(gc("Días programados pendientes")), saldo=num(gc("Saldo de días")), dvenc=num(gc("Días a vencimiento")),
                      est=txt(gc("Estado vigencia")), nat=num(gc("N.° asistencias")), notas=txt(gc("Notas"))[:500]))
hitos.sort(key=lambda h: h["f"], reverse=True)
data = dict(meta=dict(hoy=datetime.date.today().isoformat(), fuente=os.path.basename(SRC), generado=datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
                      chk=CHK_S, dias=[d for _, d in days]), ats=ats, cal=cal, contratos=contr, hitos=hitos)
json.dump(data, open(os.path.join(HERE, "data_at.json"), "w", encoding="utf-8"), ensure_ascii=False)
print("asistencias", len(ats), "| con calendario", len(cal), "| contratos", len(contr), "| hitos", len(hitos), "| fuente", SRC)
print("etapas", collections.Counter(a["etapa"] for a in ats)); print("semáforo", collections.Counter(a["sem"] for a in ats))
print("códigos calendario", collections.Counter(m[1] for v in cal.values() for m in v))
